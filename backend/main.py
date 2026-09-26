from fastapi import FastAPI, Query, Request, HTTPException, Depends, status, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse, HTMLResponse, RedirectResponse
from fastapi.security import APIKeyHeader
from starlette.middleware.base import BaseHTTPMiddleware
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
from pydantic import BaseModel
from typing import Optional
import io
import csv
import os
import secrets
import bcrypt
from datetime import datetime, timedelta
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
import logging
from logging.handlers import RotatingFileHandler

from database import (
    init_db, get_trades, update_note,
    update_exit_price, delete_all_trades,
    group_fills_into_trades, add_grouped_trade,
    add_open_position, remove_closed_open_positions,
    supabase, TABLE
)
from delta_client import fetch_all_filled_orders, fetch_open_positions

load_dotenv(dotenv_path="../.env")

app = FastAPI(title="Trading Journal", docs_url=None, redoc_url=None)

# ============================================
# AUDIT LOGGING SETUP
# ============================================

# Log file path
LOG_FILE = "/home/ubuntu/trading-journal/audit.log"

# Rotating file handler (max 10MB, 5 backups)
handler = RotatingFileHandler(
    LOG_FILE,
    maxBytes=10 * 1024 * 1024,  # 10MB
    backupCount=5
)
handler.setFormatter(logging.Formatter(
    '%(asctime)s | %(levelname)s | %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
))

logger = logging.getLogger("trading_journal")
logger.setLevel(logging.INFO)
logger.addHandler(handler)

# Console pe bhi print karo (journalctl ke liye)
console_handler = logging.StreamHandler()
console_handler.setFormatter(logging.Formatter(
    '%(asctime)s | %(levelname)s | %(message)s'
))
logger.addHandler(console_handler)


def get_client_ip(request: Request) -> str:
    """Client IP nikalta hai (proxy headers ke saath)"""
    forwarded = request.headers.get("X-Forwarded-For")
    if forwarded:
        return forwarded.split(",")[0].strip()
    real_ip = request.headers.get("X-Real-IP")
    if real_ip:
        return real_ip
    return request.client.host if request.client else "unknown"

# ============================================
# LOGIN CONFIG
# ============================================

LOGIN_USER = os.getenv("LOGIN_USER", "admin")
LOGIN_PASS_HASH = os.getenv("LOGIN_PASS_HASH", "")
API_KEY = os.getenv("API_KEY", "changeme")

# Session store (in-memory)
SESSIONS = {}
SESSION_DURATION_HOURS = 2160  # 90 days


def hash_password(password: str) -> str:
    """bcrypt hash karta hai"""
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()

def verify_login(username: str, password: str) -> bool:
    """Username/password verify karta hai (bcrypt)"""
    if not secrets.compare_digest(username, LOGIN_USER):
        return False
    try:
        return bcrypt.checkpw(password.encode(), LOGIN_PASS_HASH.encode())
    except Exception as e:
        print(f"Password verify error: {e}")
        return False


def create_session(username: str) -> str:
    """Naya session banata hai"""
    session_id = secrets.token_urlsafe(32)
    SESSIONS[session_id] = {
        "user": username,
        "expires": datetime.now() + timedelta(hours=SESSION_DURATION_HOURS)
    }
    return session_id


def verify_session(session_id: str) -> Optional[str]:
    """Session verify karta hai"""
    if not session_id or session_id not in SESSIONS:
        return None
    session = SESSIONS[session_id]
    if session["expires"] < datetime.now():
        del SESSIONS[session_id]
        return None
    return session["user"]


def get_session_from_request(request: Request) -> Optional[str]:
    """Request se session nikalta hai"""
    session_id = request.cookies.get("session_id")
    return verify_session(session_id) if session_id else None


async def require_login(request: Request):
    """Login required dependency"""
    user = get_session_from_request(request)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Login required"
        )
    return user


async def verify_api_key(api_key: str = Depends(APIKeyHeader(name="X-API-Key", auto_error=False))):
    """API Key verify karta hai (cron job ke liye)"""
    if not api_key or not secrets.compare_digest(api_key, API_KEY):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid API key",
        )
    return api_key


# ============================================
# SECURITY MIDDLEWARE
# ============================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://13.50.185.179:8000",
        "https://ip-172-31-41-231.tail0cee49.ts.net",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["*"],
)

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)


class SecurityHeadersMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["X-XSS-Protection"] = "1; mode=block"
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
        response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
        response.headers["Permissions-Policy"] = "geolocation=(), microphone=(), camera=()"
        return response

app.add_middleware(SecurityHeadersMiddleware)

# ============================================
# HTTPS REDIRECT MIDDLEWARE
# ============================================

class HTTPSRedirectMiddleware(BaseHTTPMiddleware):
    """HTTP requests ko HTTPS pe redirect karta hai"""
    async def dispatch(self, request: Request, call_next):
        forwarded_proto = request.headers.get("X-Forwarded-Proto", "http")
        host = request.headers.get("Host", "")
        
        # Agar HTTP hai aur Tailscale domain se aa raha hai
        if forwarded_proto == "http" and "tail0cee49.ts.net" in host:
            https_url = f"https://{host}{request.url.path}"
            if request.url.query:
                https_url += f"?{request.url.query}"
            return RedirectResponse(url=https_url, status_code=301)
        
        return await call_next(request)


app.add_middleware(HTTPSRedirectMiddleware)


init_db()

# ============================================
# LOGIN PAGE
# ============================================

LOGIN_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Login - Trading Journal</title>
    <style>
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body {
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            background: #0d1117;
            background-image: 
                radial-gradient(ellipse at top, rgba(31, 111, 235, 0.15) 0%, transparent 50%),
                radial-gradient(ellipse at bottom, rgba(137, 87, 229, 0.1) 0%, transparent 50%),
                linear-gradient(180deg, #0a0e27 0%, #131b3a 50%, #0d1117 100%);
            min-height: 100vh;
            display: flex;
            align-items: center;
            justify-content: center;
            padding: 20px;
            color: #c9d1d9;
        }
        .login-box {
            background: rgba(22, 27, 34, 0.75);
            backdrop-filter: blur(16px);
            -webkit-backdrop-filter: blur(16px);
            border: 1px solid rgba(88, 166, 255, 0.15);
            border-radius: 16px;
            padding: 40px 32px;
            width: 100%;
            max-width: 400px;
            box-shadow: 0 20px 60px rgba(0, 0, 0, 0.4);
            position: relative;
            overflow: hidden;
        }
        .login-box::before {
            content: '';
            position: absolute;
            top: 0;
            left: 0;
            right: 0;
            height: 1px;
            background: linear-gradient(90deg, transparent, rgba(88, 166, 255, 0.6), transparent);
        }
        h1 {
            font-size: 1.8rem;
            text-align: center;
            margin-bottom: 8px;
            background: linear-gradient(135deg, #58a6ff 0%, #56d4dd 50%, #a371f7 100%);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            background-clip: text;
            font-weight: 800;
            letter-spacing: -0.5px;
        }
        .subtitle {
            text-align: center;
            color: #8b949e;
            font-size: 12.5px;
            margin-bottom: 28px;
        }
        .form-group {
            margin-bottom: 18px;
        }
        label {
            display: block;
            font-size: 11px;
            color: #8b949e;
            text-transform: uppercase;
            letter-spacing: 0.6px;
            margin-bottom: 6px;
            font-weight: 600;
        }
        input {
            width: 100%;
            background: rgba(13, 17, 23, 0.8);
            border: 1px solid rgba(48, 54, 61, 0.6);
            border-radius: 10px;
            padding: 12px 14px;
            color: #c9d1d9;
            font-size: 14px;
            font-family: inherit;
            transition: all 0.2s ease;
        }
        input:focus {
            outline: none;
            border-color: #58a6ff;
            box-shadow: 0 0 0 3px rgba(88, 166, 255, 0.15), 0 0 20px rgba(88, 166, 255, 0.15);
        }
        button {
            width: 100%;
            background: linear-gradient(135deg, #238636 0%, #2ea043 100%);
            color: white;
            border: 1px solid rgba(63, 185, 80, 0.4);
            padding: 13px;
            border-radius: 10px;
            cursor: pointer;
            font-size: 14px;
            font-weight: 700;
            letter-spacing: 0.3px;
            transition: all 0.25s ease;
            font-family: inherit;
            margin-top: 8px;
        }
        button:hover {
            transform: translateY(-1px);
            box-shadow: 0 8px 24px rgba(63, 185, 80, 0.3);
        }
        button:active { transform: translateY(0); }
        .error {
            background: rgba(248, 81, 73, 0.1);
            border: 1px solid rgba(248, 81, 73, 0.3);
            color: #f85149;
            padding: 10px 14px;
            border-radius: 8px;
            font-size: 13px;
            margin-bottom: 18px;
            text-align: center;
        }
    </style>
</head>
<body>
    <div class="login-box">
        <h1>Trading Journal</h1>
        <div class="subtitle">Sign in to continue</div>
        {ERROR}
        <form method="POST" action="/login">
            <div class="form-group">
                <label>Username</label>
                <input type="text" name="username" required autofocus autocomplete="username">
            </div>
            <div class="form-group">
                <label>Password</label>
                <input type="password" name="password" required autocomplete="current-password">
            </div>
            <button type="submit">Sign In</button>
        </form>
    </div>
</body>
</html>
"""


@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request, error: Optional[str] = None):
    """Login page"""
    if get_session_from_request(request):
        return RedirectResponse(url="/", status_code=302)

    error_html = f'<div class="error">{error}</div>' if error else ""
    return LOGIN_HTML.replace("{ERROR}", error_html)


@app.post("/login")
@limiter.limit("3/minute")
async def login_submit(
    request: Request,
    username: str = Form(...),
    password: str = Form(...)
):
    """Login form submit"""
    client_ip = get_client_ip(request)
    user_agent = request.headers.get("User-Agent", "unknown")[:100]
    
    if verify_login(username, password):
        session_id = create_session(username)
        
        # ✅ Audit log - success
        logger.info(f"LOGIN_SUCCESS | user={username} | ip={client_ip} | ua={user_agent}")
        
        response = RedirectResponse(url="/", status_code=302)
        response.set_cookie(
            key="session_id",
            value=session_id,
            httponly=True,
            secure=False,
            samesite="lax",
            max_age=SESSION_DURATION_HOURS * 3600,
            path="/"
        )
        return response
    else:
        # ✅ Audit log - failure
        logger.warning(f"LOGIN_FAILED | user={username} | ip={client_ip} | ua={user_agent}")
        return RedirectResponse(url="/login?error=Invalid+credentials", status_code=302)


@app.get("/logout")
async def logout(request: Request):
    """Logout"""
    session_id = request.cookies.get("session_id")
    client_ip = get_client_ip(request)
    
    if session_id and session_id in SESSIONS:
        user = SESSIONS[session_id].get("user", "unknown")
        logger.info(f"LOGOUT | user={user} | ip={client_ip}")
        del SESSIONS[session_id]
    
    response = RedirectResponse(url="/login", status_code=302)
    response.delete_cookie("session_id", path="/")
    return response


# ============================================
# REQUEST MODELS
# ============================================

class NoteUpdate(BaseModel):
    note: str

class ExitUpdate(BaseModel):
    exit_price: float


# ============================================
# PUBLIC: Health check
# ============================================

@app.get("/api/health")
def health():
    return {"status": "ok"}


# ============================================
# LOGIN REQUIRED: Main endpoints
# ============================================

@app.get("/api/trades")
@limiter.limit("60/minute")
def list_trades(
    request: Request,
    date: Optional[str] = Query(None),
    user: str = Depends(require_login)
):
    all_trades = get_trades()
    if date:
        all_trades = [
            t for t in all_trades
            if (t.get("trade_time") or "").startswith(date)
        ]
    return all_trades


@app.post("/api/sync")
@limiter.limit("5/minute")
async def sync_orders(request: Request, user: str = Depends(require_login)):
    client_ip = get_client_ip(request)
    
    try:
        fills = await fetch_all_filled_orders()
        grouped = group_fills_into_trades(fills)
        removed = remove_closed_open_positions(grouped)

        synced = 0
        for trade in grouped:
            try:
                if add_grouped_trade(trade):
                    synced += 1
            except Exception as e:
                print(f"Trade skip: {e}")

        open_positions = await fetch_open_positions()
        open_count = 0
        for pos in open_positions:
            try:
                if add_open_position(pos):
                    open_count += 1
            except Exception as e:
                print(f"Open position skip: {e}")

        # ✅ Audit log
        logger.info(f"SYNC_SUCCESS | user={user} | ip={client_ip} | synced={synced} | open={open_count}")

        return {
            "synced": synced,
            "fills": len(fills),
            "trades": len(grouped),
            "open_positions": open_count,
            "removed_closed": removed
        }
    except Exception as e:
        logger.error(f"SYNC_ERROR | user={user} | ip={client_ip} | error={str(e)}")
        raise


@app.put("/api/trades/{trade_id}/note")
@limiter.limit("30/minute")
def save_note(
    request: Request,
    trade_id: int,
    body: NoteUpdate,
    user: str = Depends(require_login)
):
    update_note(trade_id, body.note)
    return {"ok": True}


@app.put("/api/trades/{trade_id}/exit")
@limiter.limit("30/minute")
async def save_exit(
    request: Request,
    trade_id: int,
    body: ExitUpdate,
    user: str = Depends(require_login)
):
    import asyncio
    loop = asyncio.get_event_loop()
    await loop.run_in_executor(None, update_exit_price, trade_id, body.exit_price)
    return {"ok": True}


# ============================================
# LOGIN REQUIRED: Exports
# ============================================

@app.get("/api/export/csv")
@limiter.limit("10/minute")
def export_csv(request: Request, date: Optional[str] = Query(None), user: str = Depends(require_login)):
    trades = get_trades()
    if date:
        trades = [t for t in trades if (t.get("trade_time") or "").startswith(date)]

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Date", "Time", "Symbol", "Side", "Size",
        "Entry Price", "Exit Price", "PnL", "Fee", "Funding", "Net PnL", "Note"
    ])

    for t in trades:
        trade_time = t.get("trade_time") or ""
        date_part = trade_time[:10] if trade_time else ""
        time_part = trade_time[11:16] if len(trade_time) > 11 else ""
        pnl = t.get("pnl", 0) or 0
        fee = t.get("fee", 0) or 0
        funding = t.get("funding", 0) or 0
        net_pnl = pnl + funding - fee if t.get("exit_price") else 0

        writer.writerow([
            date_part, time_part,
            t.get("symbol", ""), t.get("side", ""), t.get("size", ""),
            t.get("entry_price", ""), t.get("exit_price", ""),
            pnl, fee, funding, net_pnl, t.get("note", "")
        ])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=trades_{date or 'all'}.csv"}
    )


@app.get("/api/export/excel")
@limiter.limit("10/minute")
def export_excel(request: Request, date: Optional[str] = Query(None), user: str = Depends(require_login)):
    trades = get_trades()
    if date:
        trades = [t for t in trades if (t.get("trade_time") or "").startswith(date)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Trades"

    headers = [
        "Date", "Time", "Symbol", "Side", "Size",
        "Entry Price", "Exit Price", "PnL", "Fee", "Funding", "Net PnL", "Note"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="238636", end_color="238636", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for t in trades:
        trade_time = t.get("trade_time") or ""
        date_part = trade_time[:10] if trade_time else ""
        time_part = trade_time[11:16] if len(trade_time) > 11 else ""
        pnl = t.get("pnl", 0) or 0
        fee = t.get("fee", 0) or 0
        funding = t.get("funding", 0) or 0
        net_pnl = pnl + funding - fee if t.get("exit_price") else 0

        ws.append([
            date_part, time_part,
            t.get("symbol", ""), t.get("side", ""), t.get("size", ""),
            t.get("entry_price", ""), t.get("exit_price", ""),
            pnl, fee, funding, net_pnl, t.get("note", "")
        ])

    column_widths = [12, 8, 12, 8, 10, 12, 12, 12, 10, 12, 12, 40]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    net_col = 11
    for row in range(2, len(trades) + 2):
        cell = ws.cell(row=row, column=net_col)
        if cell.value and cell.value < 0:
            cell.font = Font(color="F85149", bold=True)
        elif cell.value and cell.value > 0:
            cell.font = Font(color="3FB950", bold=True)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=trades_{date or 'all'}.xlsx"}
    )


@app.get("/api/export/pdf")
@limiter.limit("10/minute")
def export_pdf(request: Request, date: Optional[str] = Query(None), user: str = Depends(require_login)):
    trades = get_trades()
    if date:
        trades = [t for t in trades if (t.get("trade_time") or "").startswith(date)]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=30, leftMargin=30,
        topMargin=30, bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle', parent=styles['Heading1'], fontSize=18,
        textColor=colors.HexColor('#1f6feb'), spaceAfter=12, alignment=1
    )
    subtitle_style = ParagraphStyle(
        'Subtitle', parent=styles['Normal'], fontSize=10,
        textColor=colors.HexColor('#8b949e'), spaceAfter=20, alignment=1
    )

    elements = []
    elements.append(Paragraph("Trading Journal", title_style))
    date_range = f"Date: {date}" if date else "All Trades"
    elements.append(Paragraph(date_range, subtitle_style))
    elements.append(Spacer(1, 10))

    closed_trades = [t for t in trades if t.get("exit_price")]
    total_pnl = sum((t.get("pnl", 0) or 0) + (t.get("funding", 0) or 0) - (t.get("fee", 0) or 0) for t in closed_trades)
    total_fee = sum(t.get("fee", 0) or 0 for t in trades)
    total_funding = sum(t.get("funding", 0) or 0 for t in trades)
    wins = sum(1 for t in closed_trades if ((t.get("pnl", 0) or 0) + (t.get("funding", 0) or 0) - (t.get("fee", 0) or 0)) > 0)
    win_rate = (wins / len(closed_trades) * 100) if closed_trades else 0

    summary_data = [
        ["Total Trades", "Closed Trades", "Net PnL", "Total Fee", "Funding", "Win Rate"],
        [str(len(trades)), str(len(closed_trades)), f"${total_pnl:.2f}", f"${total_fee:.4f}", f"${total_funding:.4f}", f"{win_rate:.1f}%"]
    ]
    summary_table = Table(summary_data, colWidths=[80, 80, 80, 80, 80, 80])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#161b22')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#f0f6fc')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#0d1117')),
        ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor('#c9d1d9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#30363d')),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))

    table_data = [["Date", "Symbol", "Side", "Size", "Entry", "Exit", "PnL", "Fee", "Net"]]
    for t in trades[:200]:
        trade_time = t.get("trade_time") or ""
        date_part = trade_time[:10] if trade_time else ""
        pnl = t.get("pnl", 0) or 0
        fee = t.get("fee", 0) or 0
        funding = t.get("funding", 0) or 0
        net_pnl = (pnl + funding - fee) if t.get("exit_price") else 0

        table_data.append([
            date_part, t.get("symbol", ""), t.get("side", "").upper(),
            str(t.get("size", "")), f"${t.get('entry_price', '')}",
            f"${t.get('exit_price', '')}" if t.get('exit_price') else "-",
            f"${pnl:.4f}", f"${fee:.4f}", f"${net_pnl:.4f}" if t.get("exit_price") else "-"
        ])

    if table_data:
        col_widths = [60, 60, 40, 30, 60, 60, 60, 60, 60]
        trade_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        trade_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f6feb')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f6f8fa')),
            ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor('#0d1117')),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#8b949e')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6f8fa')]),
        ]))
        elements.append(trade_table)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=trades_{date or 'all'}.pdf"}
    )


# ============================================
# CRON SYNC (API key)
# ============================================

@app.post("/api/cron-sync")
@limiter.limit("10/minute")
async def cron_sync(request: Request, _: str = Depends(verify_api_key)):
    return await sync_orders(request, user="cron")


# ============================================
# FRONTEND (login required)
# ============================================

@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    """Dashboard - login required"""
    if not get_session_from_request(request):
        return RedirectResponse(url="/login", status_code=302)

    with open("../frontend/index.html", "r") as f:
        return HTMLResponse(f.read())


# Static files (CSS, JS) - public
app.mount("/static", StaticFiles(directory="../frontend"), name="static")
